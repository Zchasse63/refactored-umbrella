#!/usr/bin/env python3
"""
One-time build step: extract the Yuno Group USA "RoyalStar" May-2026 US sell-sheet
(124 unique US SKUs) from the source xlsx into a normalized JSON the Portal importer
consumes. US-market (110-120V), import-ready Key Features. No images (placeholder track).

Quarantines placeholder SKUs ([TBD-*]) and emits a validation summary.

    python3 scripts/extract-yuno-us.py            # writes lib/data/source/yuno_us_appliances.json
    python3 scripts/extract-yuno-us.py --dry       # print summary only, no write
"""
import json, os, re, sys

import openpyxl

SRC = os.path.expanduser("~/Downloads/Yuno_RoyalStar_Product_Catalog_May2026.xlsx")
OUT = os.path.join(os.path.dirname(__file__), "..", "lib", "data", "source", "yuno_us_appliances.json")

HEADERS = ["Category", "Subcategory", "SKU (Recommended)", "SKU on Sell Sheet",
           "Product Name", "Capacity / Size", "Key Features", "Dimensions", "Voltage", "Status / Notes"]


def slugify(s: str) -> str:
    s = (s or "").lower().strip()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return re.sub(r"(^-|-$)", "", s)


def split_features(raw: str):
    if not raw:
        return []
    # Key Features are comma/semicolon separated clauses.
    parts = re.split(r"\s*[;,]\s*", str(raw).strip())
    out, seen = [], set()
    for p in parts:
        p = p.strip().rstrip(".")
        if p and p.lower() not in seen:
            seen.add(p.lower())
            out.append(p)
    return out


def is_us_voltage(v: str) -> bool:
    v = str(v or "")
    return "110" in v or "120" in v


def has_foreign_voltage(v: str) -> bool:
    v = str(v or "")
    return any(x in v for x in ("220", "230", "240"))


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Product Catalog"]
    rows = list(ws.iter_rows(values_only=True))

    # find header row (the one containing "Category" and "SKU (Recommended)")
    hdr_idx = next(i for i, r in enumerate(rows)
                   if r and "Category" in [str(c).strip() if c else c for c in r]
                   and any(str(c).strip() == "SKU (Recommended)" for c in r if c))
    header = [str(c).strip() if c else "" for c in rows[hdr_idx]]
    col = {h: header.index(h) for h in HEADERS if h in header}

    products, quarantined, seen_sku = [], [], set()
    cat_counts, verify = {}, []

    for r in rows[hdr_idx + 1:]:
        if not r:
            continue
        get = lambda h: (r[col[h]] if h in col and col[h] < len(r) else None)
        sku = (str(get("SKU (Recommended)")).strip() if get("SKU (Recommended)") else "")
        name = (str(get("Product Name")).strip() if get("Product Name") else "")
        category = (str(get("Category")).strip() if get("Category") else "")
        if not sku or not name or not category:
            continue
        # quarantine placeholders / malformed SKUs
        if sku.startswith("[") or "TBD" in sku.upper():
            quarantined.append({"sku": sku, "name": name, "reason": "placeholder SKU"})
            continue
        if sku in seen_sku:
            quarantined.append({"sku": sku, "name": name, "reason": "duplicate SKU (kept first)"})
            continue
        seen_sku.add(sku)

        sub = (str(get("Subcategory")).strip() if get("Subcategory") else "")
        cap = (str(get("Capacity / Size")).strip() if get("Capacity / Size") else "")
        dims = (str(get("Dimensions")).strip() if get("Dimensions") else "")
        volt = (str(get("Voltage")).strip() if get("Voltage") else "")
        feats_raw = get("Key Features")
        notes = (str(get("Status / Notes")).strip() if get("Status / Notes") else "")

        specs = []
        if cap and cap.lower() not in ("—", "-", "none"):
            specs.append({"label": "Capacity / Size", "value": cap})
        if dims and dims.lower() not in ("—", "-", "none"):
            specs.append({"label": "Dimensions", "value": dims})
        if volt and volt.lower() not in ("—", "-", "none"):
            specs.append({"label": "Voltage", "value": volt})

        needs_verify = bool(notes) and not re.search(r"confirmed", notes, re.I)
        if needs_verify:
            verify.append({"sku": sku, "note": notes})

        products.append({
            "sku": sku,
            "sku_sell_sheet": (str(get("SKU on Sell Sheet")).strip() if get("SKU on Sell Sheet") else None),
            "slug": f"{slugify(name)}-{slugify(sku)}",
            "name": name,
            "category": category,
            "subcategory": sub or None,
            "specs": specs,
            "features": split_features(feats_raw),
            # US sell sheet is 110-120V; only flag if a foreign voltage is explicitly listed
            "voltage_flag": has_foreign_voltage(volt),
            "us_voltage_confirmed": is_us_voltage(volt),
            "status_note": notes or None,
            "needs_verify": needs_verify,
        })
        cat_counts[category] = cat_counts.get(category, 0) + 1

    payload = {
        "source": "Yuno Group USA — RoyalStar US Sell-Sheet (May 2026)",
        "market": "US 110-120V",
        "count": len(products),
        "products": products,
    }

    # ---- validation summary ----
    print(f"EXTRACTED {len(products)} products  (quarantined {len(quarantined)})")
    print("\nBy category:")
    for c, n in sorted(cat_counts.items(), key=lambda x: -x[1]):
        print(f"  {n:3d}  {c}")
    print(f"\nVoltage: {sum(1 for p in products if p['us_voltage_confirmed'])} US-confirmed, "
          f"{sum(1 for p in products if p['voltage_flag'])} foreign-voltage flag, "
          f"{sum(1 for p in products if not p['us_voltage_confirmed'] and not p['voltage_flag'])} blank/unstated")
    print(f"Specs: avg {sum(len(p['specs']) for p in products)/max(len(products),1):.1f}/product; "
          f"features avg {sum(len(p['features']) for p in products)/max(len(products),1):.1f}/product")
    print(f"Needs-verify (Status/Notes non-confirmed): {len(verify)}")
    if quarantined:
        print("\nQuarantined:")
        for q in quarantined:
            print(f"  [{q['reason']}] {q['sku']} — {q['name']}")
    # slug collisions within this set
    slugs = [p["slug"] for p in products]
    dupe_slugs = {s for s in slugs if slugs.count(s) > 1}
    if dupe_slugs:
        print("\n!! DUPLICATE SLUGS:", dupe_slugs)

    if "--dry" not in sys.argv:
        with open(OUT, "w") as f:
            json.dump(payload, f, indent=2, ensure_ascii=False)
        print(f"\nWROTE {os.path.relpath(OUT)}")


if __name__ == "__main__":
    main()
